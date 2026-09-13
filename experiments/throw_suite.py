"""Create 36 Blender throws, record 108 4K streams, predict and tabulate errors."""
import argparse
import csv
from concurrent.futures import ThreadPoolExecutor, as_completed
import html
import json
from pathlib import Path
import subprocess
import sys
import time

import numpy as np

from trajectory_prediction.observations import project
from trajectory_prediction.run import load_truth

ROOT = Path(__file__).resolve().parents[1]
CAMERAS = ('Stereo_Left', 'Stereo_Right', 'Side_Overview')
LABELS = {'003_cracker_box': '饼干盒', '004_sugar_box': '糖盒', '005_tomato_soup_can': '番茄汤罐',
          '006_mustard_bottle': '芥末瓶', '007_tuna_fish_can': '金枪鱼罐', '011_banana': '香蕉'}


def save_json(path, value):
    temp = path.with_suffix('.tmp.json')
    temp.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False)+'\n')
    temp.replace(path)


def stage(command, log):
    print('  Running:', log.parent.parent.parent.name, log.parent.parent.name, log.name, flush=True)
    with log.open('w') as stream:
        result = subprocess.run(command, cwd=ROOT, stdout=stream, stderr=subprocess.STDOUT)
    if result.returncode:
        raise RuntimeError(f'Stage failed ({result.returncode}): {log}\n'+log.read_text()[-3500:])


def validate_recording(folder, config, model, preset):
    meta, records, times, truth = load_truth(folder)
    assert meta['moving_models'] == [model] and len(meta['objects']) == 1
    assert len(records) == config['frames'] and meta['fps'] == config['fps']
    np.testing.assert_allclose(meta['trajectory']['start_m'], preset['start_m'], atol=1e-6)
    np.testing.assert_allclose(meta['trajectory']['launch_velocity_m_s'], preset['velocity_m_s'], atol=1e-6)
    expected = np.array(preset['start_m']) + times[:, None]*np.array(preset['velocity_m_s'])
    expected[:, 2] -= .5*9.81*times**2
    np.testing.assert_allclose(truth, expected, atol=2e-5)
    bounds = np.array(meta['objects'][0]['bounds_local_m'])
    result = {'videos': {}, 'minimum_image_margin': {}, 'frames': len(times)}
    for name in CAMERAS:
        probe = json.loads(subprocess.check_output(['ffprobe', '-v', 'error', '-count_frames',
            '-select_streams', 'v:0', '-show_entries', 'stream=width,height,nb_read_frames,avg_frame_rate',
            '-of', 'json', str(folder / (name+'.mp4'))], text=True))['streams'][0]
        a, b = map(int, probe['avg_frame_rate'].split('/'))
        assert a/b == config['fps'] and int(probe['nb_read_frames']) == config['frames']
        assert [probe['width'], probe['height']] == [config['width'], config['height']]
        result['videos'][name] = probe
        camera = meta['cameras'][name]
        other = [json.loads(s) for s in (folder/'poses'/f'{name}.jsonl').read_text().splitlines()]
        assert len(other) == len(records)
        margin = 1.
        for source, row in zip(records, other):
            assert row['frame'] == source['frame'] and row['timestamp_s'] == source['timestamp_s']
            assert len(row['objects']) == 1
            np.testing.assert_allclose(row['objects'][0]['object_to_world'], source['objects'][0]['object_to_world'], atol=1e-8)
            world = np.array(row['objects'][0]['object_to_world'])
            np.testing.assert_allclose(np.array(camera['world_to_camera_opencv']) @ world,
                                       row['objects'][0]['object_to_camera'], atol=1e-6)
            corners = bounds @ world[:3, :3].T + world[:3, 3]
            uv = project(corners, camera) / [camera['width'], camera['height']]
            margin = min(margin, float(np.min(np.minimum(uv, 1-uv))))
        assert margin >= (.099 if name == 'Side_Overview' else .079), (name, margin)
        result['minimum_image_margin'][name] = margin
    left, right = [np.array(meta['cameras'][n]['world_to_camera_opencv']) for n in CAMERAS[:2]]
    relative = right @ np.linalg.inv(left)
    np.testing.assert_allclose(relative[:3, :3], np.eye(3), atol=1e-6)
    np.testing.assert_allclose(relative[:3, 3], [-.24, 0, 0], atol=1e-6)
    result['launch_speed_m_s'] = float(np.linalg.norm(preset['velocity_m_s']))
    result['observed_peak_speed_m_s'] = float(np.linalg.norm(np.diff(truth, axis=0)*meta['fps'], axis=1).max())
    save_json(folder/'verification.json', result)
    return result


def summarize(root, config):
    rows = []
    for model in config['models']:
        for preset in config['presets']:
            trial = root/model/preset['id']
            if not (trial/'complete.json').exists():
                continue
            report = json.loads((trial/'prediction/metrics.json').read_text())
            verified = json.loads((trial/'motion/verification.json').read_text())
            for method, metric in report['metrics'].items():
                rows.append({
                    '物体': LABELS[model], '模型': model, '组别': preset['id'], '轨迹': preset['label'],
                    '初速度_m_s': verified['launch_speed_m_s'], '峰值速度_m_s': verified['observed_peak_speed_m_s'],
                    **{f'起点_{axis}_m': v for axis, v in zip('XYZ', preset['start_m'])},
                    **{f'初速度_V{axis}_m_s': v for axis, v in zip('XYZ', preset['velocity_m_s'])},
                    '翻滚_deg_s': preset['spin_deg_s'], '方法': method, '观测模式': report['observation_mode'],
                    '观测帧': report['observed_frames'], '预测帧': report['predicted_frames'],
                    '预测时距_s': report['forecast_duration_s'], '像素噪声_px': report['noise_px'], '随机种子': report['seed'],
                    '观测ADE_cm': report['observation_ADE_m']*100,
                    '预测ADE_cm': metric['ADE_m']*100, '终点FDE_cm': metric['FDE_m']*100,
                    'RMSE3D_cm': metric['RMSE_3D_m']*100, '最大误差_cm': metric['max_error_m']*100,
                    '原始视频目录': str(trial.relative_to(root)/'motion'),
                    '预测视频目录': str(trial.relative_to(root)/'prediction/views/gravity'),
                    '预测明细CSV': str(trial.relative_to(root)/'prediction/predictions.csv')})
    if not rows:
        return
    columns = list(rows[0])
    main = [r for r in rows if r['方法'] == 'gravity']
    for filename, subset in [('summary.csv', main), ('all_methods.csv', rows)]:
        with (root/filename).open('w', encoding='utf-8-sig', newline='') as stream:
            writer = csv.DictWriter(stream, fieldnames=columns, lineterminator='\n')
            writer.writeheader()
            writer.writerows(subset)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, subset in [('重力预测_每组一行', main), ('三种方法对比', rows)]:
        sheet = workbook.create_sheet(title)
        sheet.append(columns)
        for row in subset:
            sheet.append([row[c] for c in columns])
        sheet.freeze_panes = 'E2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='17456B')
            cell.font = Font(color='FFFFFF', bold=True)
        for i, name in enumerate(columns, 1):
            sheet.column_dimensions[get_column_letter(i)].width = 25 if name in ('模型', '组别') else 19
        for cells in sheet.iter_rows(min_row=2):
            for cell in cells:
                if isinstance(cell.value, float):
                    cell.number_format = '0.0000'
        for i, row in enumerate(subset, 2):
            for name in ('原始视频目录', '预测视频目录'):
                cell = sheet.cell(i, columns.index(name)+1)
                cell.hyperlink = row[name]+'/index.html'
                cell.style = 'Hyperlink'
    sheet = workbook.create_sheet('逐帧坐标_重力预测')
    detail_columns = ['物体', '模型', '组别', '帧', '时间_s', '预测时距_s',
                      '预测X_m', '预测Y_m', '预测Z_m', '真实X_m', '真实Y_m', '真实Z_m',
                      '偏差X_m', '偏差Y_m', '偏差Z_m', '三维误差_cm', '预测Vx_m_s', '预测Vy_m_s', '预测Vz_m_s']
    sheet.append(detail_columns)
    for row in main:
        with (root/row['预测明细CSV']).open(newline='') as stream:
            for detail in csv.DictReader(stream):
                if detail['method'] != 'gravity':
                    continue
                pred = [float(detail[f'pred_{a}_m']) for a in 'xyz']
                truth = [float(detail[f'true_{a}_m']) for a in 'xyz']
                sheet.append([row['物体'], row['模型'], row['组别'], int(detail['frame']),
                              float(detail['time_s']), float(detail['horizon_s']), *pred, *truth,
                              *[a-b for a, b in zip(pred, truth)], float(detail['error_m'])*100,
                              *[float(detail[f'v{a}_m_s']) for a in 'xyz']])
    sheet.freeze_panes = 'D2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='17456B')
        cell.font = Font(color='FFFFFF', bold=True)
    for i in range(1, len(detail_columns)+1):
        sheet.column_dimensions[get_column_letter(i)].width = 22
    for cells in sheet.iter_rows(min_row=2):
        for cell in cells:
            if isinstance(cell.value, float):
                cell.number_format = '0.0000'
    sheet = workbook.create_sheet('评估说明')
    notes = [
        ['实验数量', f"已完成 {len(main)} / {len(config['models'])*len(config['presets'])} 组"],
        ['视频规格', f"{config['width']}×{config['height']}, {config['fps']} FPS, {config['frames']} 帧"],
        ['观测与预测', f"前 {config['observe_frames']} 帧历史；其余帧为未见未来，固定时刻一次预测"],
        ['观测模式', 'synthetic_stereo：真实中心投影+1px独立高斯噪声+双目三角化；没有从视频运行检测器'],
        ['目标点', '各物体几何中心，世界坐标米，Z向上；不是测量质心或翻滚中的模型原点'],
        ['真值来源', '每组 Blender pose JSONL；未来真值仅用于评估'],
        ['动力学', '理想重力自由飞行、指定翻滚；无空气阻力/碰撞，最后一帧不是实际落地'],
        ['控制变量', '六个物体使用同六组初始条件；同组共享噪声种子。物体形状不进入中心预测器。'],
        ['相机', '双目保持24cm基线与位姿，共同调焦覆盖目标；侧面逐组自动覆盖全轨迹。内外参逐组保存。'],
        ['误差', 'ADE=平均3D距离；FDE=末帧距离；RMSE3D=3D距离平方均值开方；全部只评未来'],
        ['统计限制', '每个条件一次噪声实现；不代表真实视觉精度，也不是训练后的跨物体泛化评测'],
        ['源码配置', 'experiments/throw_suite.json；根目录 manifest.json 保存本次配置'],
    ]
    for row in notes:
        sheet.append(row)
    sheet.column_dimensions['A'].width, sheet.column_dimensions['B'].width = 24, 130
    workbook.save(root/'summary.tmp.xlsx')
    (root/'summary.tmp.xlsx').replace(root/'summary.xlsx')
    body = ''.join('<tr>'+''.join(f'<td>{html.escape(str(r[c])) if not isinstance(r[c], float) else f"{r[c]:.3f}"}</td>'
                                for c in ('物体', '轨迹', '初速度_m_s', '预测ADE_cm', '终点FDE_cm'))
                   +f'<td><a href="{r["原始视频目录"]}/index.html">三路原始视频</a> · <a href="{r["预测视频目录"]}/index.html">预测视频</a></td></tr>' for r in main)
    page = f'''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>36组抛掷预测实验</title>
<style>body{{font:16px/1.7 system-ui;max-width:1300px;margin:30px auto;padding:20px;background:#f5f7fb}}table{{border-collapse:collapse;width:100%}}td,th{{padding:10px;border-bottom:1px solid #ccc}}a{{color:#1654ba}}</style>
<h1>六个物体 × 六组抛掷</h1><p>已完成 {len(main)} / 36 组。默认重力预测，前 {config['observe_frames']} 帧观测。输入为带噪声的双目仿真观测，尚未从视频检测物体。</p>
<p><a href="summary.xlsx">Excel 总表</a> · <a href="summary.csv">36组重力预测 CSV</a> · <a href="all_methods.csv">三种方法对比 CSV</a></p>
<table><tr><th>物体</th><th>轨迹</th><th>初速度 m/s</th><th>ADE cm</th><th>FDE cm</th><th>视频</th></tr>{body}</table></html>'''
    (root/'index.html').write_text(page, encoding='utf-8')



def process_trial(root, config, model, preset):
    trial = root/model/preset['id']
    trial.mkdir(parents=True, exist_ok=True)
    logs = trial/'logs'
    logs.mkdir(exist_ok=True)
    print(f'START {model} / {preset["id"]}', flush=True)
    required = [trial/'scene.blend', trial/'prediction/metrics.json', trial/'prediction/predictions.csv']
    required += [trial/k/(name+'.mp4') for k in ('motion', 'prediction/views/gravity') for name in CAMERAS]
    if (trial/'complete.json').exists() and all(p.is_file() for p in required):
        print('  Already complete', flush=True)
        return
    scene, motion, prediction = trial/'scene.blend', trial/'motion', trial/'prediction'
    launcher = [sys.executable, str(ROOT/'blender/launch.py')]
    if not (motion/'verification.json').exists() or not all((motion/(n+'.mp4')).is_file() for n in CAMERAS):
        stage(launcher + ['create-motion', '--moving-object', model, '--output', str(scene),
              '--frames', str(config['frames']), '--fps', str(config['fps']),
              '--launch-position', *map(str, preset['start_m']), '--launch-velocity', *map(str, preset['velocity_m_s']),
              '--spin-deg-s', str(preset['spin_deg_s']), '--spin-axis', *map(str, preset['spin_axis']), '--fit-stereo'], logs/'create.log')
        stage(launcher + ['record', '--scene', str(scene), '--output', str(motion), '--width', str(config['width']),
              '--height', str(config['height']), '--samples', str(config['samples'])], logs/'record.log')
        validate_recording(motion, config, model, preset)
    stage([sys.executable, '-m', 'trajectory_prediction.run', '--motion', str(motion), '--output', str(prediction),
           '--observe-frames', str(config['observe_frames']), '--noise-px', str(config['noise_px']),
           '--seed', str(preset['seed'])], logs/'predict.log')
    stage([sys.executable, '-m', 'trajectory_prediction.render_views', '--motion', str(motion),
           '--prediction', str(prediction)], logs/'overlay.log')
    # Correct the generic recorder's navigation for this nested experiment directory.
    page = motion/'index.html'
    page.write_text(page.read_text().replace('../trajectory_prediction/views/gravity/index.html',
                                            '../prediction/views/gravity/index.html'))
    save_json(trial/'complete.json', {'model': model, 'preset': preset['id'], 'status': 'complete',
                                     'observation_mode': 'synthetic_stereo'})

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--config', type=Path, default=ROOT/'experiments/throw_suite.json')
    parser.add_argument('--output', type=Path, default=ROOT/'results/throw_suite')
    parser.add_argument('--workers', type=int, choices=(1, 2, 3), default=2, help='Independent Blender jobs; default 2')
    parser.add_argument('--limit', type=int, help='Only process the first N trials (resume later without this option)')
    parser.add_argument('--summarize-only', action='store_true')
    args = parser.parse_args()
    config = json.loads(args.config.read_text())
    root = args.output.resolve()
    root.mkdir(parents=True, exist_ok=True)
    manifest = root/'manifest.json'
    if manifest.exists() and json.loads(manifest.read_text()) != config:
        parser.error('Output configuration differs. Use another --output folder.')
    save_json(manifest, config)
    if args.summarize_only:
        summarize(root, config)
        return
    started = time.monotonic()
    jobs = [(m, p) for m in config['models'] for p in config['presets']]
    if args.limit is not None:
        jobs = jobs[:args.limit]
    with ThreadPoolExecutor(max_workers=args.workers) as pool:
        pending = {pool.submit(process_trial, root, config, model, preset): (model, preset['id'])
                   for model, preset in jobs}
        for number, future in enumerate(as_completed(pending), 1):
            try:
                future.result()
            except Exception:
                for task in pending:
                    task.cancel()
                raise
            summarize(root, config)
            print(f'COMPLETE {number}/{len(jobs)} | {pending[future]} | elapsed {(time.monotonic()-started)/60:.1f} min', flush=True)
    summarize(root, config)
    print('Saved results:', root/'summary.xlsx', flush=True)


if __name__ == '__main__':
    main()

"""Verify all saved experiment results, including full decoding of overlay videos."""
import argparse
from concurrent.futures import ThreadPoolExecutor
import csv
import json
from pathlib import Path
import subprocess

import numpy as np
from openpyxl import load_workbook

from .throw_suite import ROOT, CAMERAS, save_json
from trajectory_prediction.run import load_truth


def verify_trial(job):
    root, config, model, preset = job
    trial = root/model/preset['id']
    assert (trial/'complete.json').is_file(), str(trial)
    assert (trial/'scene.blend').is_file()
    meta, records, times, truth = load_truth(trial/'motion')
    assert meta['moving_models'] == [model]
    checks = json.loads((trial/'motion/verification.json').read_text())
    assert len(checks['videos']) == 3
    report = json.loads((trial/'prediction/metrics.json').read_text())
    assert report['target'] == model+' geometry_center_world_m'
    assert report['observed_frames'] == config['observe_frames']
    assert report['seed'] == preset['seed']
    with (trial/'prediction/predictions.csv').open(newline='') as f:
        predictions = list(csv.DictReader(f))
    for method in ('gravity', 'constant_velocity', 'constant_acceleration'):
        rows = [r for r in predictions if r['method'] == method]
        count = config['observe_frames']
        assert len(rows) == config['frames']-count
        p = np.array([[float(r[f'pred_{a}_m']) for a in 'xyz'] for r in rows])
        gt = np.array([[float(r[f'true_{a}_m']) for a in 'xyz'] for r in rows])
        np.testing.assert_allclose(gt, truth[count:], atol=1e-10)
        np.testing.assert_allclose([float(r['time_s']) for r in rows], times[count:], atol=1e-10)
        assert [int(r['frame']) for r in rows] == [r['frame'] for r in records[count:]]
        error = np.linalg.norm(p-gt, axis=1)
        np.testing.assert_allclose([float(r['error_m']) for r in rows], error, atol=1e-10)
        metric = report['metrics'][method]
        np.testing.assert_allclose([metric['ADE_m'], metric['FDE_m'], metric['RMSE_3D_m']],
                                   [error.mean(), error[-1], np.sqrt(np.mean(error**2))], atol=1e-10)
    for camera in CAMERAS:
        assert (trial/'motion'/f'{camera}.mp4').is_file()
        video = trial/'prediction/views/gravity'/f'{camera}.mp4'
        probe = json.loads(subprocess.check_output(['ffprobe', '-v', 'error', '-threads', '1',
            '-count_frames', '-select_streams', 'v:0', '-show_entries',
            'stream=width,height,nb_read_frames,avg_frame_rate', '-of', 'json', str(video)], text=True))['streams'][0]
        a, b = map(int, probe['avg_frame_rate'].split('/'))
        assert a/b == config['fps'] and int(probe['nb_read_frames']) == config['frames']
        assert [probe['width'], probe['height']] == [config['width'], config['height']]
    return {'model': model, 'group': preset['id'], 'status': 'passed'}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, default=ROOT/'results/throw_suite')
    args = parser.parse_args()
    root = args.output.resolve()
    config = json.loads((root/'manifest.json').read_text())
    jobs = [(root, config, m, p) for m in config['models'] for p in config['presets']]
    results = []
    with ThreadPoolExecutor(max_workers=3) as pool:
        for result in pool.map(verify_trial, jobs):
            results.append(result)
            print(f'Checked {len(results)}/{len(jobs)}: {result["model"]} {result["group"]}', flush=True)
    expected = len(jobs)
    for filename, count in [('summary.csv', expected), ('all_methods.csv', expected*3)]:
        with (root/filename).open(encoding='utf-8-sig', newline='') as f:
            rows = list(csv.DictReader(f))
        assert len(rows) == count, (filename, len(rows))
        for row in rows:
            report = json.loads((root/row['模型']/row['组别']/'prediction/metrics.json').read_text())
            np.testing.assert_allclose(float(row['预测ADE_cm']), report['metrics'][row['方法']]['ADE_m']*100, atol=1e-8)
    book = load_workbook(root/'summary.xlsx', read_only=True)
    assert book['重力预测_每组一行'].max_row == expected+1
    assert book['三种方法对比'].max_row == expected*3+1
    assert book['逐帧坐标_重力预测'].max_row == expected*(config['frames']-config['observe_frames'])+1
    book.close()
    save_json(root/'verification.json', {'status': 'passed', 'trials': expected,
        'raw_camera_videos': expected*3, 'overlay_camera_videos': expected*3,
        'summary_rows': expected, 'method_rows': expected*3,
        'gravity_future_rows': expected*(config['frames']-config['observe_frames']), 'checks': results})
    print('All experiment videos, forecasts and spreadsheet counts verified.', flush=True)


if __name__ == '__main__':
    main()
